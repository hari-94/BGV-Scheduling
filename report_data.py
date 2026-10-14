"""
report_data.py — the numbers behind the Report page: how many rooms each RQS
inspected and each housekeeper cleaned, read from the team's own record
(GC8 Inspections, one tab a day), beside what the staff schedule had each
person down for.

Two halves, no Streamlit in either:

  * Reading. `read_workbook` takes every day tab, one per date, as room rows.
    The office PC does this at 4 PM (hotsos_agent.sync_report) and
    stores the rows here, a month to an app_settings key -- the page cannot
    see SharePoint, and the workbook is 3 MB.
  * Counting. `Directory` turns the names people type ("Danny R.", "Castor",
    "Ana Casia") into one person each, and `build` counts by person and day:
    rooms inspected and cleaned from the sheet, days on the schedule from the
    staff weeks.

What counts, and why:
  * Inspected: the first name in a room's RQS cell ("Oralia/Josseling" is
    Oralia with a trainee), on Full Clean rooms -- in-house cleans, early outs
    and deep cleans included. A Daily Service with an RQS name was looked at
    on RQS 2's round, not inspected room by room, so it is not counted (HP,
    11 Oct). Dust n Vac is counted on its own.
  * Cleaned: everybody in a room's HSKP cell -- "Nancy/Edward" counts for
    both, which is also how an RQS who cleaned rooms is seen.
  * A name in the RQS column that belongs to a manager, a houseperson or a
    housekeeper is not an inspection by an RQS; it is left out, and listed.
"""
import collections
import datetime as _dt
import hashlib
import json
import re
import unicodedata

import clock
import hotsos_sync as hs

#: Raised whenever the page needs something this module didn't have before:
#: the page reloads an older copy that a running app is still holding.
VERSION = 3
SYNC_KEY = "report_sync"              # {at, by, source, ..., checks, months: {YYYY-MM: hash}}
MONTH_PREFIX = "report_days_"         # + YYYY-MM: {"days": {iso: {"tab", "rows"}}}
REQUEST_KEY = "report_sync_request"   # {id, by, at}: the page's "Sync now"
NAMES_KEY = "report_names"            # {name as typed on the sheet: person}: the page's own fixes
LEFT_KEY = "staff_left"               # {person: {"by", "at"}}: no longer works here (marked on the page)
SYNC_HOUR = 16                        # 4 PM, property time: the day's work is in

# ── service families ─────────────────────────────────────────────────────────
FC, DC, DS, DV = "FC", "DC", "DS", "DV"   # full clean, deep clean, daily service, dust n vac
INSPECTED = (FC, DC)                      # what an RQS inspects room by room


def family(service) -> str:
    """The job a Service cell names, or "" for none we count (P/U models,
    touch-ups, blanks). "Full Clean (IH)", "E/O 10/03" and "CHECK OUT" are
    Full Cleans; "Stayover" and the exchanges are Daily Service."""
    s = " ".join(str(service or "").lower().split())
    if not s:
        return ""
    if "dust" in s:
        return DV
    if "deep" in s:
        return DC
    if s.startswith(("daily", "sheet", "stay", "towel")):
        return DS
    if "p/u" in s or "model" in s:
        return ""
    if "full" in s or s.startswith(("check out", "checkout", "early out", "e/o", "eo ")):
        return FC
    return ""


# ── reading the workbook ─────────────────────────────────────────────────────
_ROOM = re.compile(r"^[1-9]\d{3}[A-I]$")
# Tabs the app itself left behind to compare its charts with the team's.
_APP_COPY = re.compile(r"\b(app|output)\b|_app_|app_result", re.I)


def _label(v) -> str:
    return " ".join(str(v if v is not None else "").lower().split())


def _col_kind(label):
    """Which column a header label names. The team has typed the
    housekeeper column HSKP, HSK, HSKPNG, HKP, Housekeeping and Housekeeper."""
    if not label:
        return None
    if label.startswith("room") or label == "unit":
        return "room"
    if label.startswith("service") or label == "clean type":
        return "service"
    if label.startswith(("hsk", "hkp", "housekeep", "attendant", "room attendant")) or label == "hk":
        return "hskp"
    if label.startswith(("rq", "inspector")):
        return "rqs"
    return None


# Headers of the columns that are none of ours -- never taken for a missing one.
_OTHER_LABELS = ("time", "pet", "current guest", "res type", "notes", "notas", "status",
                 "late", "arriving", "carpet", "stripping", "stripper")


def _share(rows, j, test, n=60):
    """The share of the filled cells of column j, in the first n rows, that
    pass `test`."""
    vals = [str(r[j]).strip() for r in rows[:n] if j < len(r) and str(r[j] or "").strip()]
    return sum(1 for v in vals if test(v)) / len(vals) if vals else 0.0


def read_tab(rows):
    """(room rows, repeats removed, notes) from one day tab, given its rows
    as tuples.

    The header is the row in the first few that names most of Room, Service,
    HSKP and RQS. Several tabs have lost a label, and a missing column is
    found by what it holds; `notes` says which:
      "Sep 19"  "p" where "Room" was  -> the column of room numbers
      "July 19" no "Service"          -> the column of Full Clean / Daily Service
      "Aug 8", "JAN 5"  no "RQS" (blank; "Hold, COD/RENTAL STATUS")
                                      -> the column of names just after HSKP
      "Oct 1"   no "HSKP"             -> the column of names just before RQS"""
    rows = [tuple(r) for r in rows]
    best, start = None, 0
    for i, r in enumerate(rows[:8]):
        found = {}
        for j, v in enumerate(r):
            k = _col_kind(_label(v))
            if k and k not in found:
                found[k] = j
        if len(found) >= 2 and (best is None or len(found) > len(best)):
            best, start = found, i + 1
        if best and len(best) == 4:
            break
    if best is None:
        raise ValueError("no header row")
    cols, notes = dict(best), []
    header = rows[start - 1]
    body = rows[start:]
    if "room" not in cols:
        hits = collections.Counter(j for r in body[:60] for j, v in enumerate(r[:4])
                                   if _ROOM.match(_label(v).upper()))
        if not hits:
            raise ValueError("no Room column")
        cols["room"] = hits.most_common(1)[0][0]
        notes.append("Room column found by its room numbers")
    rooms = [r for r in body if cols["room"] < len(r) and _ROOM.match(_label(r[cols["room"]]).upper())]
    width = max([len(r) for r in rooms[:60]] + [0])

    def free(j):
        lab = _label(header[j] if j < len(header) else "")
        return (0 <= j < width and j not in cols.values() and _col_kind(lab) is None
                and not lab.startswith(_OTHER_LABELS))

    def names_in(j):
        return _share(rooms, j, lambda v: is_name(v) and not re.search(r"\d", v)) >= 0.6

    if "service" not in cols:
        scored = sorted(((_share(rooms, j, lambda v: family(v) != ""), j)
                         for j in range(width) if free(j)), reverse=True)
        if scored and scored[0][0] >= 0.6:
            cols["service"] = scored[0][1]
            notes.append("Service column found by what's in it")
    if "rqs" not in cols and "hskp" in cols and free(cols["hskp"] + 1) and names_in(cols["hskp"] + 1):
        cols["rqs"] = cols["hskp"] + 1
        notes.append("RQS column found next to the housekeepers")
    if "hskp" not in cols and "rqs" in cols and free(cols["rqs"] - 1) and names_in(cols["rqs"] - 1):
        cols["hskp"] = cols["rqs"] - 1
        notes.append("housekeeper column found next to the RQS")
    if "service" not in cols or not ({"rqs", "hskp"} & set(cols)):
        raise ValueError("no Service, HSKP or RQS column")

    def cell(r, key):
        j = cols.get(key)
        return " ".join(str(r[j]).split()) if j is not None and j < len(r) and r[j] is not None else ""

    out, seen, repeats = [], {}, 0
    for r in body:
        room = cell(r, "room").upper()
        if not _ROOM.match(room):
            continue
        fam = family(cell(r, "service"))
        row = [room, fam, cell(r, "hskp"), cell(r, "rqs")]
        key = (room, fam)
        if key in seen:                      # the same room and job twice
            repeats += 1
            old = out[seen[key]]
            if sum(map(bool, row[2:])) > sum(map(bool, old[2:])):
                out[seen[key]] = row         # keep the line that was filled in
            continue
        seen[key] = len(out)
        out.append(row)
    return out, repeats, notes


def _filled(rows):
    return sum(1 for r in rows if r[2] and r[3])


def read_workbook(src, year=None, upto=None):
    """Every day tab of the workbook at `src` (a path or a file object).

    Returns {"days": {iso: {"tab", "rows"}}, "checks": {...}}. One tab per
    date: where two claim a date, the one with more rooms filled in wins,
    then the original over a "(2)" copy. Days after `upto` are left out --
    a tab built ahead for tomorrow is a plan, not work done."""
    import openpyxl
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    try:
        year = year or clock.today().year
        by_date, no_date, app_copies, unreadable, repeats = {}, [], [], [], 0
        order = []
        for name in wb.sheetnames:
            if _APP_COPY.search(name.replace("_", " ")) or "_app_" in name.lower():
                app_copies.append(name)
            else:
                order.append([name, hs.tab_date(name, year)])
        redated = place_dates(order)
        found = []
        for name, d in order:
            if d is None or d.year != year:
                no_date.append(name)
                continue
            if upto and d > upto:
                continue
            try:
                rows, rep, notes = read_tab(wb[name].iter_rows(values_only=True))
            except ValueError as ex:
                unreadable.append(f"{name}: {ex}")
                continue
            repeats += rep
            found += [f"{name}: {n}" for n in notes]
            by_date.setdefault(d, []).append((name, rows))
    finally:
        wb.close()
    days, doubles = {}, {}
    for d, tabs in by_date.items():
        tabs.sort(key=lambda t: (_filled(t[1]), "(2)" not in t[0]), reverse=True)
        days[d.isoformat()] = {"tab": tabs[0][0], "rows": tabs[0][1]}
        if len(tabs) > 1:
            doubles[d.isoformat()] = {"used": tabs[0][0], "not used": [t[0] for t in tabs[1:]]}
    # A tab copied to make the next day and never changed carries yesterday's
    # names under today's date. None so far; worth saying if one appears.
    sig, copies = {}, []
    for iso in sorted(days):
        rows = days[iso]["rows"]
        if _filled(rows) < 10:
            continue
        h = hashlib.md5(json.dumps(sorted(rows)).encode()).hexdigest()
        if h in sig:
            copies.append(f"{days[iso]['tab']} is the same as {days[sig[h]]['tab']}")
        else:
            sig[h] = iso
    have = sorted(days)
    missing = []
    if have:
        d = _dt.date.fromisoformat(have[0])
        last = _dt.date.fromisoformat(have[-1])
        while d <= last:
            if d.isoformat() not in days:
                missing.append(d.isoformat())
            d += _dt.timedelta(days=1)
    return {"days": days,
            "checks": {"tabs_used": len(days), "two_tabs": doubles, "app_copies": app_copies,
                       "other_tabs": no_date, "unreadable": unreadable,
                       "repeats_removed": repeats, "copied_tabs": copies,
                       "missing_dates": missing, "redated": redated,
                       "columns_found": found,
                       "first": have[0] if have else None, "last": have[-1] if have else None}}


def place_dates(order) -> dict:
    """Date the tabs whose name says one thing and whose place another.

    The tabs run in date order, so a tab's neighbours say what day it can
    be: "Jun 26" between "jan 25" and "Jan27" is Jan 26 (a typo for "Jan"),
    and "9" between "April 8th" and "APRIL 10" is Apr 9. Only where the
    neighbours leave exactly one day with that day number; a tab in its
    place, or with neighbours far apart, keeps its name's date. `order` is
    [[name, date or None], ...] and is corrected in place; returns
    {name: ISO date} for the tabs it re-dated."""
    out = {}
    for i, (name, d) in enumerate(order):
        prev = next((order[k][1] for k in range(i - 1, -1, -1) if order[k][1]), None)
        nxt = next((order[k][1] for k in range(i + 1, len(order)) if order[k][1]), None)
        if not prev or not nxt or not 2 <= (nxt - prev).days <= 4:
            continue
        if d and prev <= d <= nxt:
            continue
        bare = re.fullmatch(r"\s*(\d{1,2})\s*", name)
        day = d.day if d else (int(bare.group(1)) if bare else None)
        if day is None:
            continue
        fits = [prev + _dt.timedelta(days=k) for k in range(1, (nxt - prev).days)
                if (prev + _dt.timedelta(days=k)).day == day]
        if len(fits) == 1:
            order[i][1] = fits[0]
            out[name] = fits[0].isoformat()
    return out




# ── names ────────────────────────────────────────────────────────────────────
_NOISE = {"bld", "bldg", "building", "lead", "am", "pm", "gtl", "b", "hsp", "rqs",
          "w", "with", "con", "x", "cama", "bed", "bano"}
_VISITING = re.compile(r"\b(peak\s*7|gl\s*7|glt|gtl)\b", re.I)
# Words that begin a note rather than a name: "Ventanas", "Already
# inspected", "Need Housekeeper 1", "Salen mañana", "In House".
_NOT_A_NAME = {
    "none", "no", "manager", "managers", "unassigned", "tbd", "verify", "not", "na",
    "checked", "check", "clean", "cleaning", "ok", "dnd", "dust", "inspector",
    "inspected", "hk", "hskp", "daily", "vacant", "ventanas", "windows", "pisos",
    "floors", "deep", "training", "houseperson", "hspm", "hspam", "pending", "blocked",
    "should", "salen", "already", "yesterday", "tomorrow", "today", "early", "late",
    "in", "house", "need", "needs", "help", "original", "it", "done", "out", "room",
    "stripping", "carpets", "carpet", "linen", "eng", "engineering", "front", "desk",
    "model", "models", "rqs", "same", "see", "call", "sick", "off", "kit", "kitchen",
    "cocina", "retoques", "touch", "cleaned", "occupied", "spotlight", "yesteday", "self"}


def _ascii(s) -> str:
    return unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()


def words(s) -> list:
    """A name as lowercase words, punctuation and tags dropped:
    "SAYRA. H" -> ["sayra", "h"], "Willy - BLD 2" -> ["willy"],
    "Mayra w/" -> ["mayra"]."""
    w = re.findall(r"[a-z]+", _ascii(s).lower())
    return [x for i, x in enumerate(w) if i == 0 or x not in _NOISE]


def _key(s) -> str:
    return " ".join(words(s))


def is_name(s) -> bool:
    w = words(s)
    return bool(w) and len(w[0]) >= 2 and w[0] not in _NOT_A_NAME \
        and not _VISITING.search(str(s))


def _osa(a, b) -> int:
    """Edit distance with a swap of two neighbours as one edit ("Casia")."""
    d = [[0] * (len(b) + 1) for _ in range(len(a) + 1)]
    for i in range(len(a) + 1):
        d[i][0] = i
    for j in range(len(b) + 1):
        d[0][j] = j
    for i in range(1, len(a) + 1):
        for j in range(1, len(b) + 1):
            c = 0 if a[i - 1] == b[j - 1] else 1
            d[i][j] = min(d[i - 1][j] + 1, d[i][j - 1] + 1, d[i - 1][j - 1] + c)
            if i > 1 and j > 1 and a[i - 1] == b[j - 2] and a[i - 2] == b[j - 1]:
                d[i][j] = min(d[i][j], d[i - 2][j - 2] + 1)
    return d[len(a)][len(b)]


def _abbrev(r, w) -> bool:
    """"hdz" for Hernandez, "rdgz" for Rodriguez: same first letter, and the
    rest in order."""
    if not r or not w or r[0] != w[0]:
        return False
    it = iter(w)
    return all(c in it for c in r)


def _rest_fits(r, later) -> bool:
    """Does `r`, typed after a first name, fit one of the real later names?"""
    for w in later:
        if len(r) <= 2:
            if w.startswith(r):
                return True
        elif w.startswith(r) or r.startswith(w) or _abbrev(r, w) or \
                (r[0] == w[0] and _osa(r, w) <= max(1, len(w) // 4)):
            return True
    return False


def _first_fits(t, w) -> int:
    """How well a typed first name fits a real one: 100 the same; 75 a short
    form or a slip at the end ("Edwar", "Ale", "Gustavoyy"); 70 a typo
    ("Byran", "Camile", "Elizebeth"); 0 none. A short name one letter off is
    a different name, not a typo: Celia is not Delia."""
    if t == w:
        return 100
    if len(t) >= 3 and w.startswith(t):
        return 75
    if len(w) >= 4 and t.startswith(w) and len(t) - len(w) <= 2:
        return 75
    if len(t) >= 4 and t[0] == w[0] and abs(len(t) - len(w)) <= 2:
        d = _osa(t, w)
        if d == 1 and (len(w) >= 6 or len(t) != len(w) or sorted(t) == sorted(w)):
            return 70
        if d == 2 and len(w) >= 8:
            return 70
    return 0


def _score(tw, nw) -> int:
    if tw == nw:
        return 100
    f = _first_fits(tw[0], nw[0])
    if f and (len(tw) == 1 or _rest_fits(tw[1], nw[1:])):
        return 90 if f == 100 else f
    if len(tw) >= 2 and len(tw[-1]) >= 4 and tw[-1] in nw[1:] and nw[0][0] == tw[0][0]:
        return 80                       # "JENNI CAICEDO": the surname is the sure part
    if len(tw) <= 2 and len(tw[0]) >= 4 and tw[0] in nw[1:] \
            and (len(tw) == 1 or nw[0].startswith(tw[1])):
        return 60                       # "Castor", "Urbina", "Castor M"
    if len(tw) >= 2 and len(tw[0]) >= 3 and nw[0][:3] == tw[0][:3] and _rest_fits(tw[1], nw[1:]):
        return 50                       # "Jeni C", "Jenny S"
    return 0


def best(typed, names):
    """([names], score): every name in `names` that `typed` fits best."""
    tw = words(typed)
    out, top = [], 0
    if not tw:
        return out, top
    for n in names:
        nw = words(n)
        s = _score(tw, nw) if nw else 0
        if s > top:
            out, top = [n], s
        elif s and s == top:
            out.append(n)
    return out, top


def match(typed, names) -> list:
    return best(typed, names)[0]


_ARROW = re.compile(r"\s*(?:-+>|>+|→|=>)\s*")


def members(cell, first_only=False) -> list:
    """The people in a cell: "Jenifer S/ Ana C" -> ["Jenifer S", "Ana C"].

    In the RQS column only the first counts -- "Oralia/Josseling" is Oralia
    with a trainee. A room handed on, "Gloria R>>Yadira", is the last
    person's. Notes and help from a sister property are not people."""
    s = " ".join(str(cell or "").split())
    s = re.sub(r"\([^)]*(\)|$)", " ", s)              # "Jose M (Bed", "(Bryan-Ceci)"
    s = _ARROW.split(s)[-1]
    out = []
    for p in re.split(r"\s*(?:/+|,|&|\+|\band\b|\by\b)\s*|\s*[-–]\s+|\s+[-–]\s*", s):
        p = p.strip(" .*\"'#-")
        if p and is_name(p):
            out.append(p)
            if first_only:
                break
    return out


class Directory:
    """One name per person.

    The people are the staff schedule's, under HotSOS's full name where
    HotSOS has one (staff_names keeps that list). A name typed on the sheet
    is matched to them: the name tables first -- this page's own fixes, then
    the approved table the HotSOS push uses, then the roster's confirmed
    spellings -- then only matches that can't be wrong. Where a name could
    be two people, the one the column is about wins (an RQS in the RQS
    column), then the one the schedule had at work that day."""

    def __init__(self, roster_names=(), attendants=(), aliases=None, extra=None, merges=None):
        self.attendants = sorted({a for a in attendants if a})
        raw_alias = {}
        for table in (merges or {}), (aliases or {}), (extra or {}):     # later wins
            for k, v in table.items():
                if _key(k) and v:
                    raw_alias[_key(k)] = v
        # One spelling per schedule name, case and full stops aside.
        spell = {}
        for n in roster_names:
            c = self.clean(n)
            if c and is_name(c):
                spell.setdefault(_key(c), collections.Counter())[c] += 1
        plain = {k: max(v, key=lambda s: (s != s.upper(), len(s), v[s])) for k, v in spell.items()}

        # The tables name some people by a short form ("jorge luis", "nancy"):
        # point each entry at a real spelling before anything uses it.
        def target(v):
            if v in self.attendants:
                return v
            hit, score = best(v, self.attendants)
            if len(hit) == 1 and score >= 70:
                return hit[0]
            return plain.get(_key(v)) or v
        self.alias = {k: target(v) for k, v in raw_alias.items()}
        # A name with a surname or an initial: HotSOS's full name for it...
        full = {}
        for k, n in plain.items():
            if len(k.split()) > 1:
                full[k] = self._named(n, self.attendants)
        # ...or a longer spelling of the same person on the schedule: "JORGE L".
        for k, n in plain.items():
            if len(k.split()) > 1 and not full[k]:
                others = [plain[o] for o in plain
                          if o != k and len(o.split()) > 1 and not full[o] and len(o) > len(k)]
                hit, score = best(n, others)
                if len(hit) == 1 and score >= 70:
                    full[k] = hit[0]
        # A bare first name stands for the one person who has it, HotSOS and
        # the schedule both counted: "Jennifer" is nobody while "Jennifer C"
        # is Jennyfer Caicedo and HotSOS has a Jennifer of its own.
        multi = {k: (full[k] or plain[k]) for k in full}
        for k, n in plain.items():
            if len(k.split()) == 1:
                if k in self.alias:
                    full[k] = self.alias[k]
                    continue
                owners = {a for a in self.attendants if _key(a).split()[:1] == [k]}
                owners |= {v for m, v in multi.items() if m.split()[0] == k}
                full[k] = owners.pop() if len(owners) == 1 else None
        self._by_key = {k: (full.get(k) or plain[k]) for k in plain}
        self.names = sorted(set(self._by_key.values()))
        # The tables may name a person by a short form ("jorge luis"): point
        # each at the person it means.
        self.alias = {k: self._canon(v) for k, v in self.alias.items()}
        self._hits = {}

    @staticmethod
    def clean(name) -> str:
        """A schedule name without its section or building: "Willy - BLD 2"."""
        base = str(name).split(" · ")[0]
        base = re.split(r"\s+[-–]\s*(?=bld|b\d|lead)|\s+(?=(?:lead\s+)?bldg?\b)", base, flags=re.I)[0]
        return " ".join(base.split()).strip(" .-")

    def _named(self, n, pool, floor=70):
        k = _key(n)
        if k in self.alias:
            return self.alias[k]
        hit, score = best(n, pool)
        return hit[0] if len(hit) == 1 and score >= floor else None

    def _canon(self, v):
        if v in self.names:
            return v
        k = _key(v)
        if k in self._by_key:
            return self._by_key[k]
        hit, score = best(v, self.names)
        return hit[0] if len(hit) == 1 and score >= 70 else v

    def person(self, name) -> str:
        """The person a schedule name is (its own name if nobody else's)."""
        c = self.clean(name)
        return self._by_key.get(_key(c)) or c

    def who(self, typed, prefer=(), working=()):
        """The person a name typed on the sheet means, or None.

        `prefer`: the people the column is about. `working`: whom the
        schedule had at work that day. Both only settle a tie."""
        if not is_name(typed):
            return None
        k = _key(typed)
        if k in self.alias:
            return self.alias[k]
        if k not in self._hits:
            hit, _ = best(typed, self.names)
            if not hit and len(k.split()) == 1:
                # "DIANISS": one letter off a name the tables already know.
                hit = sorted({self.alias[a] for a in self.alias
                              if len(a.split()) == 1 and _first_fits(k, a) >= 70})
            self._hits[k] = hit
        hit = self._hits[k]
        for pool in (prefer, working):
            if len(hit) <= 1:
                break
            narrowed = [n for n in hit if n in pool]
            if narrowed:
                hit = narrowed
        return hit[0] if len(hit) == 1 else None


# ── the staff schedule ───────────────────────────────────────────────────────
WORKED = ("working", "daily_service", "other_duty")      # roster_import.WORKED_KINDS


def roster_days(weeks, overrides, start, end):
    """Every staff-schedule cell from `start` to `end` (ISO dates), as
    ({name as written: {iso: [(group, section, kind, cell), ...]}},
     {names of help from another property})."""
    import roster_import as ri
    s0, e0 = _dt.date.fromisoformat(start), _dt.date.fromisoformat(end)
    out, visiting = {}, set()
    for wk in sorted(weeks):
        try:
            w0 = _dt.date.fromisoformat(wk)
        except ValueError:
            continue
        if w0 + _dt.timedelta(days=6) < s0 or w0 > e0:
            continue
        eff, _ = ri.apply_overrides(weeks[wk], overrides or {}, wk)
        dates = [d for d in (eff.get("dates") or []) if start <= d <= end]
        for key, rec in (eff.get("people") or {}).items():
            name, home = ri.read_person(key, rec)
            if not name:
                continue
            if home:
                visiting.add(name)
            mine = out.setdefault(name, {})
            for iso in dates:
                kind, _ = ri.cell_kind(rec, iso)
                cell = " ".join(str((rec.get("cells") or {}).get(iso, "") or "").split())
                mine.setdefault(iso, []).append(
                    (rec.get("group") or "other", rec.get("section") or "", kind, cell))
    return out, visiting


def worked_span(roster, directory) -> dict:
    """{person: (first, last)}: the first and last day the staff schedule has
    them working. Given a roster that runs past today, someone leaving shows
    their last shift; someone who started in July isn't averaged over
    January."""
    out = {}
    for name, by_day in roster.items():
        days = [iso for iso, entries in by_day.items() if any(e[2] in WORKED for e in entries)]
        if days and is_name(Directory.clean(name)):
            p = directory.person(name)
            a, b = out.get(p, (min(days), max(days)))
            out[p] = (min(a, min(days)), max(b, max(days)))
    return out


def active(span, start, end):
    """start..end cut down to the part someone was on the schedule, or None
    if they weren't on it at all then."""
    if not span:
        return start, end
    a, b = max(start, span[0]), min(end, span[1])
    return (a, b) if a <= b else None


# ── who did what ─────────────────────────────────────────────────────────────
def _bump(d, key, name=None):
    if name is None:
        d[key] = d.get(key, 0) + 1
    else:
        d.setdefault(key, {})
        d[key][name] = d[key].get(name, 0) + 1


def build(days, roster, visiting, directory):
    """Who did what, a person and a day at a time.

    `days` is the synced sheet ({iso: {"tab", "rows"}}); `roster` and
    `visiting` come from roster_days. Returns

      people  {person: {"rqs", "hk", "other": days worked in each, "visiting",
                        "role": "rqs" | "hk" | "other" | "visiting"}}
      days    {person: {iso: {"in", "dv", "fc", "ds", "dvc": [rooms],
                              "sched": [(group, section, kind, cell)]}}}
      loose   {iso: what could not be put to an RQS or a housekeeper}

    A person's role is the team the schedule had them in most, the sheet
    agreeing: Hari is an RQS who sometimes cleans. A team row alone isn't
    enough -- Erick's HSP route numbers sat in the RQS section for 37 days
    of houseperson work, and he never inspected a room."""
    people, per, working = {}, {}, {}
    for name, by_day in roster.items():
        if not is_name(Directory.clean(name)):         # "GTL", "PEAK 7", "Carpets"
            continue
        p = directory.person(name)
        info = people.setdefault(p, {"rqs": 0, "hk": 0, "other": 0, "visiting": True, "_d": {}})
        if name not in visiting:
            info["visiting"] = False
        for iso, entries in by_day.items():
            slot = per.setdefault(p, {}).setdefault(iso, {})
            slot.setdefault("sched", []).extend(entries)
            for g, _sec, kind, _cell in entries:
                if kind in WORKED:
                    info["_d"].setdefault(g if g in ("rqs", "hk") else "other", set()).add(iso)
                    working.setdefault(iso, set()).add(p)
    for p, info in people.items():
        for g, ds in info.pop("_d").items():
            info[g] = len(ds)
    rqs = {p for p, i in people.items() if i["rqs"] and not i["visiting"]}
    hk = {p for p, i in people.items() if i["hk"] and not i["visiting"]}
    team = rqs | hk

    def put(p, iso, key, room):
        per.setdefault(p, {}).setdefault(iso, {}).setdefault(key, []).append(room)

    loose = {}
    for iso, day in days.items():
        at_work = working.get(iso, set())
        lo = {}
        for room, fam, hskp, rqs_cell in day["rows"]:
            if fam in INSPECTED or fam == DV:
                names = members(rqs_cell, first_only=True)
                who = directory.who(names[0], rqs, at_work) if names else None
                if who in rqs:
                    put(who, iso, "in" if fam in INSPECTED else "dv", room)
                elif fam in INSPECTED:
                    if not names:
                        _bump(lo, "in_none")
                    elif who:
                        _bump(lo, "in_other", who)
                    else:
                        _bump(lo, "in_unknown", names[0])
            if fam in (FC, DC, DS, DV):
                key = {FC: "fc", DC: "fc", DS: "ds", DV: "dvc"}[fam]
                names = members(hskp)
                if not names and fam != DV:
                    _bump(lo, "cl_none")
                for n in names:
                    who = directory.who(n, hk, at_work)
                    if who in team:
                        put(who, iso, key, room)
                    elif fam != DV:
                        if who:
                            vis = (people.get(who) or {}).get("visiting")
                            _bump(lo, "cl_visiting" if vis else "cl_other", who)
                        else:
                            _bump(lo, "cl_unknown", n)
        if lo:
            loose[iso] = lo
    for p, info in people.items():
        sheet = per.get(p, {})
        n_in = sum(len(s.get("in", ())) for s in sheet.values())
        n_cl = sum(len(s.get("fc", ())) + len(s.get("ds", ())) for s in sheet.values())
        can_rqs = info["rqs"] and (n_in or info["rqs"] >= info["other"])
        can_hk = info["hk"] and (n_cl or info["hk"] >= info["other"])
        info["role"] = ("visiting" if info["visiting"] else
                        "rqs" if can_rqs and (not can_hk or info["rqs"] >= info["hk"]) else
                        "hk" if can_hk else "other")
        if info["role"] in ("rqs", "hk"):
            continue
        # Not an RQS or a housekeeper after all: their rooms go with the
        # other staff's, where the page lists them.
        for iso, s in sheet.items():
            for key, lk in (("in", "in_other"), ("fc", "cl_other"), ("ds", "cl_other")):
                for _ in s.pop(key, ()):
                    _bump(loose.setdefault(iso, {}), lk, p)
            s.pop("dv", None)
            s.pop("dvc", None)
    return {"people": people, "days": per, "loose": loose}


# ── a period ─────────────────────────────────────────────────────────────────
_DOW = ["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]
# A project cell is a sentence ("Rollaways 8-9am and then, Todo el PASILLO BLD
# 2"); the card names the duty.
_DUTY = [(r"rollaway", "Rollaways"), (r"\beng\b|engineering", "ENG"), (r"\bullr\b", "ULLR"),
         (r"project|proyect", "Projects"), (r"garage", "Garages"),
         (r"\b\d{4}[a-i]\b", "Set rooms"),
         (r"\bfood\b|comida|cocina|kitchen", "Food / kitchen"), (r"stripping|linen", "Stripping linen"),
         (r"\bhsp\b|houseperson", "HSP cover"), (r"training|entrenamiento", "Training"),
         (r"carpet|alfombra", "Carpets"), (r"window|ventana", "Windows"),
         (r"deep\s*clean", "Deep cleans"), (r"breck\s*in", "Breck Inn"), (r"\bmaps?\b", "Maps"),
         (r"inventor", "Inventory"), (r"meeting|conference|session", "Meetings"),
         (r"\bsales\b", "Sales"), (r"pasillo|hallway|stairs|escalera|baseboard", "Hallways / stairs"),
         (r"touch\s*up|retoque", "Touch-ups"), (r"\blead\b", "Lead")]


def duty(cell) -> str:
    """What a duty cell is, in a word or two."""
    s = str(cell or "").lower()
    for rx, label in _DUTY:
        if re.search(rx, s):
            return label
    return " ".join(str(cell or "").split()[:3]) or "Other duty"


def tally(by_day, start, end, role):
    """One person's numbers over start..end (ISO), as `role` ("rqs"/"hk")."""
    import roster_import as ri
    t = collections.Counter()
    dows, projects, teams = collections.Counter(), collections.Counter(), collections.Counter()
    for iso, slot in by_day.items():
        if not (start <= iso <= end):
            continue
        for k in ("in", "dv", "fc", "ds", "dvc"):
            t[k] += len(slot.get(k, ()))
        if slot.get("in"):
            t["days_in"] += 1
        if slot.get("fc") or slot.get("ds"):
            t["days_clean"] += 1
        entries = slot.get("sched", ())
        mine = [e for e in entries if e[0] == role]
        worked = [e for e in mine if e[2] in WORKED]
        kinds = {e[2] for e in worked}
        if worked:
            t["sched"] += 1
            dows[_DOW[(_dt.date.fromisoformat(iso).weekday() + 1) % 7]] += 1
            if "other_duty" in kinds and not kinds & {"working", "daily_service"}:
                t["project"] += 1
                projects[" / ".join(sorted({duty(e[3]) for e in worked if e[3]})) or "Other duty"] += 1
            elif role == "hk" and "daily_service" in kinds and "working" not in kinds:
                t["ds_days"] += 1
            else:
                t["main"] += 1
            # The roles the schedule names: "RQS 1", "RQS 2", or "RQS 1 + RQS 2"
            # (one inspector doing both, counted for both). Most RQS days only
            # say ON -- an ordinary Full Clean inspecting day.
            roles = set().union(*(ri.rqs_roles(e[3]) for e in worked)) if role == "rqs" else set()
            if roles:
                t["role_days"] += 1
            if 1 in roles:
                t["rqs1"] += 1
            if 2 in roles:
                t["rqs2"] += 1
        elif any(e[2] == "no_call" for e in entries):
            t["nocall"] += 1
        elif any(e[2] == "vto" for e in mine):
            t["vto"] += 1
        if not worked:
            other = [e for e in entries if e[0] != role and e[2] in WORKED]
            if other:
                t["other_team"] += 1
                teams[other[0][1] or other[0][0]] += 1
    t["cleaned"] = t["fc"] + t["ds"]
    out = dict(t)
    out["per_day"] = round(t["in"] / t["days_in"], 1) if t["days_in"] else None
    out["dows"], out["projects"], out["teams"] = dict(dows), dict(projects), dict(teams)
    return out


def usual_days(dows) -> list:
    """The days of the week someone works most -- 60% of their busiest day
    or more, as the roster's Attendance tab has it."""
    if not dows:
        return []
    top = max(dows.values())
    return [d for d in _DOW if dows.get(d, 0) >= max(1, top * 0.6)]


def weeks_in(start, end) -> float:
    return ((_dt.date.fromisoformat(end) - _dt.date.fromisoformat(start)).days + 1) / 7


def rows_for(res, start, end, role):
    """The table for `role` over start..end: a dict per person with anything
    to show, best first."""
    out = []
    for p, by_day in res["days"].items():
        if (res["people"].get(p) or {}).get("role") != role:
            continue
        t = tally(by_day, start, end, role)
        if not (t.get("sched") or t.get("in") or t.get("cleaned") or t.get("nocall")):
            continue
        out.append({"person": p, **t})
    key = "in" if role == "rqs" else "cleaned"
    out.sort(key=lambda r: (-r.get(key, 0), r["person"]))
    rank, prev = 0, None
    for i, r in enumerate(out, 1):
        if r.get(key, 0) != prev:
            rank, prev = i, r.get(key, 0)
        r["rank"] = rank
    return out


def slices(start, end):
    """How to break a period into cards: months for anything longer than
    five weeks, weeks for a month, days for a week. [(label, start, end)]."""
    s, e = _dt.date.fromisoformat(start), _dt.date.fromisoformat(end)
    span = (e - s).days + 1
    out = []
    if span > 35:
        d = s.replace(day=1)
        while d <= e:
            nxt = (d.replace(day=28) + _dt.timedelta(days=4)).replace(day=1)
            a, b = max(d, s), min(nxt - _dt.timedelta(days=1), e)
            out.append((f"{d:%b}" + (f" {d:%Y}" if s.year != e.year else ""), a.isoformat(), b.isoformat()))
            d = nxt
    elif span > 7:
        d = s - _dt.timedelta(days=(s.weekday() + 1) % 7)
        while d <= e:
            a, b = max(d, s), min(d + _dt.timedelta(days=6), e)
            out.append((f"{a:%b} {a.day}–{b.day}" if a.month == b.month else f"{a:%b} {a.day}–{b:%b} {b.day}",
                        a.isoformat(), b.isoformat()))
            d += _dt.timedelta(days=7)
    else:
        d = s
        while d <= e:
            out.append((f"{d:%a} {d.day}", d.isoformat(), d.isoformat()))
            d += _dt.timedelta(days=1)
    return out


def loose_for(res, start, end) -> dict:
    """What could not be put to anyone over start..end, totalled."""
    out = {}
    for iso, lo in res["loose"].items():
        if start <= iso <= end:
            for k, v in lo.items():
                if isinstance(v, dict):
                    tgt = out.setdefault(k, {})
                    for n, c in v.items():
                        tgt[n] = tgt.get(n, 0) + c
                else:
                    out[k] = out.get(k, 0) + v
    return out


# ── storage ──────────────────────────────────────────────────────────────────
def workbook_year(path) -> int:
    m = re.search(r"(20\d\d)", str(path))
    return int(m.group(1)) if m else clock.today().year


def _keep(write, *args, tries=4):
    """A write that waits and asks again when the connection drops. A sync on
    the busy office PC reads for minutes, and on 14 Oct Supabase had closed
    the idle connection by the time it saved: "Server disconnected" twice,
    past db's own single retry, and the whole read was lost."""
    import time
    for i in range(tries):
        try:
            return write(*args)
        except Exception:
            if i == tries - 1:
                raise
            time.sleep(3 * (i + 1))


def store(db, result, stamp, by, source, saved_at=None, saved_by=""):
    """Keep the synced days, a month to a key, writing only the months that
    changed. Returns the sync record."""
    old = db._load_key(SYNC_KEY) or {}
    hashes = dict(old.get("months") or {})
    months = {}
    for iso, day in result["days"].items():
        months.setdefault(iso[:7], {})[iso] = day
    years = {m[:4] for m in months}
    written = []
    for m, mdays in sorted(months.items()):
        blob = {"days": mdays}
        h = hashlib.md5(json.dumps(blob, sort_keys=True).encode()).hexdigest()
        if hashes.get(m) != h:
            _keep(db._upsert_key, MONTH_PREFIX + m, blob)
            written.append(m)
        hashes[m] = h
    for m in list(hashes):               # a month of this year that lost all its tabs
        if m[:4] in years and m not in months:
            _keep(db._delete_key, MONTH_PREFIX + m)
            hashes.pop(m)
    rec = {"at": stamp, "by": by, "status": "done", "source": source,
           "file_saved_at": saved_at, "file_by": saved_by,
           "checks": result["checks"], "months": hashes, "written": written}
    _keep(db._upsert_key, SYNC_KEY, rec)
    return rec


def load_days(db) -> dict:
    """Every synced day: {iso: {"tab", "rows"}}."""
    out = {}
    for r in db._like_keys(MONTH_PREFIX, with_payload=True):
        p = r.get("payload")
        p = json.loads(p) if isinstance(p, str) else (p or {})
        out.update(p.get("days") or {})
    return out


# ── a file to download ───────────────────────────────────────────────────────
def to_excel(sheets, notes=()) -> bytes:
    """An .xlsx of `sheets` ({title: (columns, rows)}, rows as lists) with a
    Notes sheet; headers bold, columns sized, the header row frozen."""
    import io
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head = PatternFill("solid", fgColor="E8EEF6")
    for title, (cols, rows) in sheets.items():
        ws = wb.create_sheet(title[:31])
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True), head
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append(list(r))
        ws.freeze_panes = "A2"
        for i, col in enumerate(cols, 1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:500] if i - 1 < len(r)])
            ws.column_dimensions[get_column_letter(i)].width = min(max(8, width + 2), 60)
    if notes:
        ws = wb.create_sheet("Notes")
        for n in notes:
            ws.append([n])
        ws.column_dimensions["A"].width = 120
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
