"""
hotsos_sync.py — turn the day's sheet into HotSOS room assignments.

The flow, end to end:

  RQS presses "Preview" or "Push to HotSOS" on the HotSOS page
    -> the page writes a request to app_settings (REQUEST_KEY)
    -> hotsos_agent.py, running on the office PC, sees it within a minute
    -> reads today's tab of the inspections workbook (synced from SharePoint)
    -> reads HotSOS's rooms and attendants, builds a plan, and (for a push)
       sends one AssignRoom call per attendant
    -> writes the plan and the outcome to RESULT_KEY, which the page shows.

The agent has to be on the office PC because the two sources are inside the
company: SSRS signs in with the PC's Windows login, and the SharePoint file is
an organisation-only link, read from its OneDrive-synced copy.

This module is the part with no network and no Streamlit: picking the day's
tab, reading it, matching names, and building the plan. That is so it can be
tested against real workbooks on its own.
"""
import datetime as _dt
import re

# ── app_settings keys shared by the page and the agent ──────────────────────
REQUEST_KEY = "hotsos_push_request"     # {id, date, mode, by, at}
RESULT_KEY = "hotsos_push_result"       # {id, date, mode, status, plan, ...}
NAMES_KEY = "hotsos_names"              # {sheet name: HotSOS attendant label}
FORECAST_KEY = "ssrs_forecast"          # {pulled_at, days: [...], warnings}
FORECAST_REQUEST_KEY = "ssrs_forecast_request"   # {id, by, at}
HEARTBEAT_KEY = "hotsos_agent_heartbeat"         # {at, host}
ROSTER_STATUS_KEY = "roster_sync_status"         # last Schedule.xlsx import

# What a row of the plan can say. The page colours by these.
ASSIGN = "assign"            # will be (or was) assigned
MOVE = "move"                # assigned in HotSOS to someone else; the sheet wins
ALREADY = "already"          # HotSOS already has it with the right person
NO_ROOM = "room not in HotSOS"
NO_PERSON = "name not matched"
NO_HSKP = "no housekeeper on sheet"
SKIPPED = (NO_ROOM, NO_PERSON, NO_HSKP)

_ROOM_RE = re.compile(r"^[1-9]\d{3}[A-Z]$")
_MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct",
     "nov", "dec"], 1)}

# Names on the sheet that mean "nobody" rather than a person.
_NOT_A_PERSON = {"", "none", "-", "—", "no hk", "no hk available", "manager",
                 "unassigned", "tbd", "verify"}


def _clean(v) -> str:
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def _norm(v) -> str:
    return re.sub(r"[^a-z ]", "", _clean(v).lower()).strip()


# ── which tab is today's ─────────────────────────────────────────────────────
def tab_date(name: str, year: int):
    """The date a tab name stands for, or None.

    The tabs are named by hand, so this takes the forms people actually type:
    "10-08", "10.8", "10-08-26", "2026-10-08", "Oct 8", "8 October",
    "Thu 10-08". Month comes first for numbers, as it does in the US.
    """
    s = _clean(name).lower()
    m = re.search(r"\b(20\d\d)[-._ ](\d{1,2})[-._ ](\d{1,2})\b", s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    else:
        m = re.search(r"\b(\d{1,2})[-._ ](\d{1,2})(?:[-._ ](\d{2,4}))?\b", s)
        if m:
            mo, d = int(m.group(1)), int(m.group(2))
            y = int(m.group(3)) if m.group(3) else year
            y = y + 2000 if y < 100 else y
        else:
            mname = re.search(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?", s)
            dnum = re.search(r"\b(\d{1,2})(?:st|nd|rd|th)?\b", s)
            if not (mname and dnum):
                return None
            mo, d, y = _MONTHS[mname.group(1)], int(dnum.group(1)), year
            ynum = re.search(r"\b(20\d\d)\b", s)
            if ynum:
                y = int(ynum.group(1))
    try:
        return _dt.date(y, mo, d)
    except ValueError:
        return None


def pick_tab(sheet_names, day: _dt.date):
    """The tab for `day`, or None. If two tabs claim the day, the last wins --
    a copied tab is usually added after the one it was copied from."""
    hit = None
    for name in sheet_names:
        if tab_date(name, day.year) == day:
            hit = name
    return hit


# ── reading a tab ────────────────────────────────────────────────────────────
_HDR = {
    "room": ("room", "room #", "room number", "unit"),
    "service": ("service", "service type", "clean type"),
    "hskp": ("hskp", "housekeeper", "attendant", "room attendant", "hk"),
    "rqs": ("rqs", "inspector"),
    "time": ("time (min)", "time", "minutes", "credits"),
}


def read_rows(rows):
    """Room lines from a tab, given its rows as tuples of cell values.

    The header row is found by its labels rather than its position, because
    the tab is edited by hand. Reading stops at nothing in particular: every
    row whose Room cell is a real room code is a room, and anything else --
    blank lines, the staff summary under the schedule, notes -- is ignored.
    """
    rows = [tuple(r) for r in rows]
    cols = None
    start = 0
    for i, r in enumerate(rows[:30]):
        labels = [_clean(c).lower() for c in r]
        found = {}
        for key, names in _HDR.items():
            for j, lab in enumerate(labels):
                if lab in names and key not in found:
                    found[key] = j
        if "room" in found and "hskp" in found:
            cols, start = found, i + 1
            break
    if cols is None:
        raise ValueError("Couldn't find a header row with both a 'Room' and an "
                         "'HSKP' (or 'Housekeeper') column in the first 30 rows.")

    def cell(r, key):
        j = cols.get(key)
        return r[j] if j is not None and j < len(r) else None

    out, seen = [], set()
    for r in rows[start:]:
        room = _clean(cell(r, "room")).upper()
        if not _ROOM_RE.match(room):
            continue
        if room in seen:          # the export repeats a room on a second line
            continue
        seen.add(room)
        out.append({"room": room,
                    "service": _clean(cell(r, "service")),
                    "hskp": _clean(cell(r, "hskp")),
                    "rqs": _clean(cell(r, "rqs"))})
    return out


def read_workbook(path, day: _dt.date):
    """(tab name, room rows) for `day` from the workbook at `path`."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        tab = pick_tab(wb.sheetnames, day)
        if tab is None:
            raise LookupError(
                f"No tab for {day:%a %b %d} in {len(wb.sheetnames)} tabs. "
                f"Last few: {', '.join(wb.sheetnames[-5:])}")
        return tab, read_rows(wb[tab].iter_rows(values_only=True))
    finally:
        wb.close()


# ── names on the sheet -> HotSOS attendants ─────────────────────────────────
def match_name(sheet_name: str, attendants, saved=None):
    """The HotSOS attendant (dict with 'id' and 'label') for a sheet name.

    `saved` is the RQS's own table (NAMES_KEY) and always wins. After that,
    only matches that cannot be wrong: the full name, a unique first name, or
    a unique first name plus last initial ("Danny R."). Anything ambiguous is
    left unmatched, so the page asks rather than guessing.
    """
    by_label = {_norm(a["label"]): a for a in attendants}
    if saved and sheet_name in saved:
        return by_label.get(_norm(saved[sheet_name]))
    n = _norm(sheet_name)
    if not n:
        return None
    if n in by_label:
        return by_label[n]
    parts = n.split()
    first = parts[0]
    initial = parts[1][0] if len(parts) > 1 and len(parts[1]) == 1 else None
    hits = []
    for a in attendants:
        ap = _norm(a["label"]).split()
        if not ap or ap[0] != first:
            continue
        if initial and not (len(ap) > 1 and ap[-1].startswith(initial)):
            continue
        hits.append(a)
    return hits[0] if len(hits) == 1 else None


# ── the plan ─────────────────────────────────────────────────────────────────
def build_plan(sheet_rows, hotsos_rooms, attendants, saved_names=None):
    """What the push will do, one line per sheet room.

    `hotsos_rooms` maps room code -> {"gid": roomGlobalId, "service": str,
    "assigned_to": attendant label or ""}. Rooms already with the right
    person are left alone; rooms with somebody else are moved, because the
    sheet is the RQS's final word.
    """
    plan = []
    for r in sheet_rows:
        line = dict(r)
        # Dust n Vac is RQS 2's own round and carries no housekeeper, so the
        # RQS on the line is the person who does it.
        who = r["hskp"]
        if _norm(who) in _NOT_A_PERSON and _norm(r["service"]).startswith("dust"):
            who = r.get("rqs", "")
        line["who"] = who
        hs = hotsos_rooms.get(r["room"])
        line["hotsos_service"] = (hs or {}).get("service", "")
        line["current"] = (hs or {}).get("assigned_to", "")
        person = None
        if _norm(who) in _NOT_A_PERSON:
            line["action"] = NO_HSKP
        elif hs is None:
            line["action"] = NO_ROOM
        else:
            person = match_name(who, attendants, saved_names)
            if person is None:
                line["action"] = NO_PERSON
            elif _norm(line["current"]) == _norm(person["label"]):
                line["action"] = ALREADY
            elif line["current"]:
                line["action"] = MOVE
            else:
                line["action"] = ASSIGN
        line["hotsos_name"] = person["label"] if person else ""
        line["person_id"] = person["id"] if person else ""
        line["gid"] = (hs or {}).get("gid", "")
        plan.append(line)
    return plan


def batches(plan):
    """{person_id: [roomGlobalId, ...]} for the lines that need a call."""
    out = {}
    for line in plan:
        if line["action"] in (ASSIGN, MOVE):
            out.setdefault(line["person_id"], []).append(line["gid"])
    return out


def summary(plan):
    """Counts by action, and the sheet names that still need matching."""
    counts = {}
    for line in plan:
        counts[line["action"]] = counts.get(line["action"], 0) + 1
    unmatched = sorted({l["who"] for l in plan if l["action"] == NO_PERSON})
    return {"counts": counts, "unmatched_names": unmatched}
